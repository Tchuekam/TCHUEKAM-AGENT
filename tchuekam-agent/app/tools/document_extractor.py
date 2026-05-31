import os
import logging

logger = logging.getLogger(__name__)

def extract_text(filepath: str) -> str:
    """Extracts text from binary documents. Returns empty string on failure."""
    ext = os.path.splitext(filepath)[1].lower()
    text = ""
    try:
        if ext == ".pdf":
            text = _extract_pdf(filepath)
        elif ext == ".docx":
            text = _extract_docx(filepath)
        elif ext == ".xlsx":
            text = _extract_xlsx(filepath)
    except Exception as e:
        logger.debug(f"Failed to extract text from {filepath}: {e}")
    return text.strip()

def _extract_pdf(filepath: str) -> str:
    try:
        import fitz
    except ImportError:
        logger.warning("PyMuPDF (fitz) is not installed.")
        return ""
    text_parts = []
    try:
        with fitz.open(filepath) as doc:
            for page in doc:
                text_parts.append(page.get_text())
    except Exception as e:
        logger.debug(f"fitz extract error for {filepath}: {e}")
    return "\n".join(text_parts)

def _extract_docx(filepath: str) -> str:
    try:
        import docx
    except ImportError:
        logger.warning("python-docx is not installed.")
        return ""
    try:
        doc = docx.Document(filepath)
        return "\n".join([para.text for para in doc.paragraphs])
    except Exception as e:
        logger.debug(f"docx extract error for {filepath}: {e}")
        return ""

def _extract_xlsx(filepath: str) -> str:
    try:
        import openpyxl
    except ImportError:
        logger.warning("openpyxl is not installed.")
        return ""
    text_parts = []
    try:
        # read_only=True is faster and uses less memory
        wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                for cell in row:
                    if cell is not None:
                        text_parts.append(str(cell))
        wb.close()
    except Exception as e:
        logger.debug(f"xlsx extract error for {filepath}: {e}")
    return " ".join(text_parts)
